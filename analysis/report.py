"""
Generate an Excel workbook from picks_log.json.

Sheets:
  1. Picks Log  — one row per pick, all fields
  2. Daily Summary  — P&L and hit-rate aggregated by game date
  3. Player Summary — aggregated by player name
  4. Stat Summary   — aggregated by stat type
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Optional

_ROOT      = Path(__file__).parent.parent
_PICKS_LOG = _ROOT / "results" / "picks_log.json"
_REPORT    = _ROOT / "results" / "picks_report.xlsx"


def _load_picks() -> list[dict]:
    if not _PICKS_LOG.exists():
        return []
    try:
        with open(_PICKS_LOG) as f:
            return json.load(f)
    except Exception:
        return []


def _pnl(pick: dict) -> Optional[float]:
    contracts = pick.get("suggested_contracts", 0)
    outcome   = pick.get("outcome")
    if not contracts or outcome not in ("win", "loss"):
        return None
    cost = (pick.get("yes_ask", 50)) / 100.0
    if outcome == "win":
        return contracts * (1.0 - cost)
    return -contracts * cost


def generate_report(output_path: Path = _REPORT) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required: pip install openpyxl")

    picks = _load_picks()
    wb    = openpyxl.Workbook()

    # ── Styles ──────────────────────────────────────────────────────────────
    header_font    = Font(bold=True, color="FFFFFF")
    header_fill    = PatternFill("solid", fgColor="1F4E79")   # dark blue
    win_fill       = PatternFill("solid", fgColor="C6EFCE")   # light green
    loss_fill      = PatternFill("solid", fgColor="FFC7CE")   # light red
    pending_fill   = PatternFill("solid", fgColor="FFEB9C")   # light yellow
    center         = Alignment(horizontal="center")

    def _header_row(ws, cols: list[str]):
        for col_idx, text in enumerate(cols, 1):
            cell = ws.cell(row=1, column=col_idx, value=text)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center

    def _auto_width(ws, min_w=10, max_w=40):
        for col_cells in ws.columns:
            width = max(
                len(str(cell.value or "")) for cell in col_cells
            )
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = (
                min(max(width + 2, min_w), max_w)
            )

    # ── Sheet 1: Picks Log ───────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Picks Log"
    cols1 = [
        "Game Date", "Player", "Stat", "Line", "Side",
        "Our %", "Kalshi %", "Edge %", "Confidence",
        "Contracts", "Yes Ask (¢)", "Actual", "Outcome",
        "P&L ($)", "Resolved Date", "Ticker",
    ]
    _header_row(ws1, cols1)
    ws1.freeze_panes = "A2"

    for pick in sorted(picks, key=lambda p: (p.get("game_date", ""), p.get("player_name", ""))):
        outcome = pick.get("outcome")
        pnl_val = _pnl(pick)
        row = [
            pick.get("game_date", ""),
            pick.get("player_name", ""),
            (pick.get("stat_type", "")).upper(),
            pick.get("line"),
            pick.get("side", "YES").upper(),
            round(pick["our_probability"] * 100, 1) if pick.get("our_probability") is not None else "",
            round(pick["implied_probability"] * 100, 1) if pick.get("implied_probability") is not None else "",
            round(pick["edge"] * 100, 1) if pick.get("edge") is not None else "",
            (pick.get("confidence", "") or "").upper(),
            pick.get("suggested_contracts", 0) or 0,
            pick.get("yes_ask", ""),
            pick.get("actual_stat"),
            (outcome or "pending").upper(),
            round(pnl_val, 2) if pnl_val is not None else "",
            pick.get("resolved_date", ""),
            pick.get("ticker", ""),
        ]
        r = ws1.max_row + 1
        for c, val in enumerate(row, 1):
            ws1.cell(row=r, column=c, value=val)

        # Color code outcome column (column 13)
        outcome_cell = ws1.cell(row=r, column=13)
        if outcome == "win":
            outcome_cell.fill = win_fill
        elif outcome == "loss":
            outcome_cell.fill = loss_fill
        elif outcome is None:
            outcome_cell.fill = pending_fill

    _auto_width(ws1)

    # ── Sheet 2: Daily Summary ───────────────────────────────────────────────
    ws2 = wb.create_sheet("Daily Summary")
    cols2 = ["Date", "Bets", "Wins", "Losses", "No Data", "Pending",
             "Hit Rate", "Total P&L ($)", "Avg P&L/Bet ($)"]
    _header_row(ws2, cols2)
    ws2.freeze_panes = "A2"

    daily: dict[str, dict] = defaultdict(lambda: {
        "bets": 0, "wins": 0, "losses": 0, "no_data": 0, "pending": 0, "pnl": 0.0
    })
    for p in picks:
        d     = p.get("game_date", "?")
        daily[d]["bets"] += 1
        outcome = p.get("outcome")
        if outcome == "win":
            daily[d]["wins"] += 1
        elif outcome == "loss":
            daily[d]["losses"] += 1
        elif outcome == "no_data":
            daily[d]["no_data"] += 1
        else:
            daily[d]["pending"] += 1
        pnl_val = _pnl(p)
        if pnl_val is not None:
            daily[d]["pnl"] += pnl_val

    for d, s in sorted(daily.items()):
        resolved = s["wins"] + s["losses"]
        hit_rate = round(s["wins"] / resolved, 3) if resolved else ""
        avg_pnl  = round(s["pnl"] / resolved, 2) if resolved else ""
        row = [d, s["bets"], s["wins"], s["losses"], s["no_data"], s["pending"],
               hit_rate, round(s["pnl"], 2), avg_pnl]
        r = ws2.max_row + 1
        for c, val in enumerate(row, 1):
            ws2.cell(row=r, column=c, value=val)
        # Color P&L cell
        pnl_cell = ws2.cell(row=r, column=8)
        if isinstance(s["pnl"], (int, float)):
            pnl_cell.fill = win_fill if s["pnl"] >= 0 else loss_fill

    _auto_width(ws2)

    # ── Sheet 3: Player Summary ──────────────────────────────────────────────
    ws3 = wb.create_sheet("Player Summary")
    cols3 = ["Player", "Bets", "Wins", "Losses", "Hit Rate", "Total P&L ($)"]
    _header_row(ws3, cols3)
    ws3.freeze_panes = "A2"

    player_stats: dict[str, dict] = defaultdict(lambda: {
        "bets": 0, "wins": 0, "losses": 0, "pnl": 0.0
    })
    for p in picks:
        name    = p.get("player_name", "?")
        outcome = p.get("outcome")
        player_stats[name]["bets"] += 1
        if outcome == "win":
            player_stats[name]["wins"] += 1
        elif outcome == "loss":
            player_stats[name]["losses"] += 1
        pnl_val = _pnl(p)
        if pnl_val is not None:
            player_stats[name]["pnl"] += pnl_val

    for name, s in sorted(player_stats.items(), key=lambda x: -x[1]["bets"]):
        resolved = s["wins"] + s["losses"]
        hit_rate = round(s["wins"] / resolved, 3) if resolved else ""
        row = [name, s["bets"], s["wins"], s["losses"], hit_rate, round(s["pnl"], 2)]
        r = ws3.max_row + 1
        for c, val in enumerate(row, 1):
            ws3.cell(row=r, column=c, value=val)

    _auto_width(ws3)

    # ── Sheet 4: Stat Summary ────────────────────────────────────────────────
    ws4 = wb.create_sheet("Stat Summary")
    cols4 = ["Stat Type", "Bets", "Wins", "Losses", "Hit Rate", "Total P&L ($)"]
    _header_row(ws4, cols4)
    ws4.freeze_panes = "A2"

    stat_stats: dict[str, dict] = defaultdict(lambda: {
        "bets": 0, "wins": 0, "losses": 0, "pnl": 0.0
    })
    for p in picks:
        stat    = (p.get("stat_type", "?") or "?").upper()
        outcome = p.get("outcome")
        stat_stats[stat]["bets"] += 1
        if outcome == "win":
            stat_stats[stat]["wins"] += 1
        elif outcome == "loss":
            stat_stats[stat]["losses"] += 1
        pnl_val = _pnl(p)
        if pnl_val is not None:
            stat_stats[stat]["pnl"] += pnl_val

    for stat, s in sorted(stat_stats.items()):
        resolved = s["wins"] + s["losses"]
        hit_rate = round(s["wins"] / resolved, 3) if resolved else ""
        row = [stat, s["bets"], s["wins"], s["losses"], hit_rate, round(s["pnl"], 2)]
        r = ws4.max_row + 1
        for c, val in enumerate(row, 1):
            ws4.cell(row=r, column=c, value=val)

    _auto_width(ws4)

    # ── Save ─────────────────────────────────────────────────────────────────
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[report] Saved {output_path} ({len(picks)} picks)")
    return output_path
